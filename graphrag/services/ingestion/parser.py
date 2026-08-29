"""`ParsedDocument` / `DocumentParser`. See BLUEPRINT §6.1.

Supports exactly the four mime types in `config/base.yaml`'s `limits.allowed_upload_mimetypes`:
PDF (pymupdf), DOCX (python-docx), and plain-text/markdown (decoded directly). No dependency on
`unstructured` (ARCHITECTURE §2.3's heavier alternative) — pymupdf + python-docx is the "leaner
image" option that document explicitly names, and it's what `pyproject.toml` now pins.
"""

from __future__ import annotations

import io
from typing import Any, Final

import pymupdf
from docx import Document as _DocxDocument
from pydantic import BaseModel, ConfigDict

from graphrag.core.errors import ValidationError
from graphrag.services.ingestion.normalizer import normalize_display

PDF_MIME: Final[str] = "application/pdf"
TEXT_MIME: Final[str] = "text/plain"
MARKDOWN_MIME: Final[str] = "text/markdown"
DOCX_MIME: Final[str] = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_SUPPORTED_MIME_TYPES: Final[frozenset[str]] = frozenset(
    {PDF_MIME, TEXT_MIME, MARKDOWN_MIME, DOCX_MIME}
)

# Separator inserted between pages/paragraphs when reassembling into one `text` blob. It is
# never inside a page's own (char_start, char_end) range, so `chunk_document`'s recursive
# paragraph split (on blank lines) sees real structure without corrupting page offsets.
_JOIN_SEP: Final[str] = "\n\n"


class ParsedDocument(BaseModel):
    """Contract: `text` is the full document body; `page_offsets[i]` (0-based list index) is
    the (char_start, char_end) range of PAGE i+1 within `text`. Empty for formats without a
    native page concept (DOCX, plain text, markdown)."""

    model_config = ConfigDict(frozen=True)

    text: str
    title: str | None
    page_offsets: list[tuple[int, int]]
    mime_type: str


class DocumentParser:
    """Contract:
    - parse(bytes, mime_type) -> ParsedDocument.
    - Supports the mime types in limits.allowed_upload_mimetypes.
    - page_offsets maps page number -> (char_start, char_end) in `text`; empty for
      formats without pages.
    - Raises ValidationError for unsupported/corrupt input. Never raises library errors.
    """

    def parse(self, raw: bytes, mime_type: str) -> ParsedDocument:
        if mime_type not in _SUPPORTED_MIME_TYPES:
            raise ValidationError(
                f"unsupported mime type: {mime_type!r}", details={"mime_type": mime_type}
            )
        if mime_type == PDF_MIME:
            parsed = self._parse_pdf(raw)
        elif mime_type == DOCX_MIME:
            parsed = self._parse_docx(raw)
        else:
            parsed = self._parse_text(raw, mime_type)

        if not parsed.text.strip():
            raise ValidationError(
                "document has no extractable text", details={"mime_type": mime_type}
            )
        return parsed

    def _parse_pdf(self, raw: bytes) -> ParsedDocument:
        # pymupdf ships incomplete type stubs (untyped `open`/`close`, no Iterable protocol on
        # `Document`) — narrow `type: ignore`s below, not a blanket suppression.
        try:
            doc = pymupdf.open(stream=raw, filetype="pdf")  # type: ignore[no-untyped-call]
        except Exception as exc:
            raise ValidationError("corrupt or unreadable PDF") from exc
        try:
            if doc.page_count == 0:
                raise ValidationError("PDF has no pages")
            parts: list[str] = []
            offsets: list[tuple[int, int]] = []
            pos = 0
            page: Any
            for i, page in enumerate(doc):  # type: ignore[arg-type]
                if i > 0:
                    parts.append(_JOIN_SEP)
                    pos += len(_JOIN_SEP)
                page_text = normalize_display(page.get_text())
                start = pos
                parts.append(page_text)
                pos += len(page_text)
                offsets.append((start, pos))
            title = (doc.metadata or {}).get("title") or None
            return ParsedDocument(
                text="".join(parts), title=title, page_offsets=offsets, mime_type=PDF_MIME
            )
        except ValidationError:
            raise
        except Exception as exc:
            raise ValidationError("corrupt or unreadable PDF") from exc
        finally:
            doc.close()  # type: ignore[no-untyped-call]

    def _parse_docx(self, raw: bytes) -> ParsedDocument:
        try:
            doc = _DocxDocument(io.BytesIO(raw))
            paragraphs = [p.text for p in doc.paragraphs]
            title = doc.core_properties.title or None
        except Exception as exc:
            raise ValidationError("corrupt or unreadable DOCX") from exc
        text = normalize_display(_JOIN_SEP.join(paragraphs))
        return ParsedDocument(text=text, title=title, page_offsets=[], mime_type=DOCX_MIME)

    def _parse_text(self, raw: bytes, mime_type: str) -> ParsedDocument:
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise ValidationError("file is not valid UTF-8 text") from exc
        return ParsedDocument(
            text=normalize_display(text), title=None, page_offsets=[], mime_type=mime_type
        )
